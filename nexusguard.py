from pathlib import Path
from datetime import datetime
from urllib.parse import parse_qs, urlsplit
from playwright.sync_api import sync_playwright, expect
import openpyxl
from copy import copy
from openpyxl.styles import Alignment
import json, os
import re

def parse_attack(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    header = re.fullmatch(r"Alert ID:\s*(\S+)\s+(.+?)\s+-\s+(.+)", lines[0] if lines else "")
    if header is None:
        raise ValueError("Cannot parse attack header: expected Alert ID and start/end times")

    attack = {
        "alert_id": header.group(1),
        "start_time": header.group(2),
        "end_time": None if header.group(3) == "(ongoing)" else header.group(3),
        "tags": None,
        "bandwidth": None,
        "packet_rate": None,
    }
    labels = {
        "Status": "status",
        "Mode": "mode",
        "Type": "type",
        "Top Attack Type": "top_attack_type",
        "Duration": "duration",
        "IP": "ip",
        "Profile": "profile",
        "Max Severity Percent": "max_severity_percent",
        "Attacked Hosts": "attacked_hosts",
    }
    for key in labels.values():
        attack[key] = None

    for index, line in enumerate(lines[1:], start=1):
        if line.startswith("Attack /") or line == "Event Tags":
            attack["tags"] = line
        elif re.fullmatch(r"\d+(?:\.\d+)?\s*[KMGT]?bps", line):
            attack["bandwidth"] = line
        elif re.fullmatch(r"\d+(?:\.\d+)?\s*[KMGT]?pps", line):
            attack["packet_rate"] = line
        elif ":" in line:
            label, value = line.split(":", 1)
            if label in labels:
                value = value.strip()
                # Some labels have their value on the following line.
                if not value and index + 1 < len(lines):
                    next_line = lines[index + 1]
                    if next_line.split(":", 1)[0] not in labels:
                        value = next_line
                attack[labels[label]] = value or None

    return attack

def search_verified_attack(browser):
    state_file = Path(__file__).resolve().with_name("nexusguard_state.json")
    if state_file.exists():
        context = browser.new_context(storage_state=str(state_file))
        print("Loaded saved session; checking whether it is still valid.")
    else:
        context = browser.new_context()
    page = context.new_page()
    page.goto("https://support.nexusguard.com")

    #skip 2FA
    skip = page.get_by_text("Skip", exact=True).filter(visible=True).first
    cmc = page.get_by_text("CMC_Global_Offload", exact=True).filter(visible=True).first

    skip.or_(cmc).first.wait_for(state="visible", timeout=180000)
    print("Skip matches:", skip.count())
    if skip.is_visible():
        skip.click()

    cmc.wait_for(state="visible", timeout=15000)
    context.storage_state(path=str(state_file))
    print("Authenticated session saved")
    print("CMC matches:", cmc.count())
    cmc.click()

    alert = page.get_by_text("DDoS Alert", exact=True).first
    alert.wait_for(state="visible", timeout=15000)
    print("DDOS alert matches: ", alert.count())
    alert.click()


    page.locator("i.iconfont.nx-search.cursor-pointer:visible").click()
    attack = page.locator(".opddos-cascader-li.second-arrow").filter(
        has=page.locator(".icon-is-attack.material-icons")
    )
    attack.wait_for(state="attached", timeout=15000)
    page.get_by_role("button", name="All Tags", exact=True).click()
    attack.filter(visible=True).click(timeout=15000)

    verified = page.get_by_text("Verified Attack", exact=True).filter(visible=True)
    verified.click(timeout=15000)
    selected_tag = page.get_by_role("button", name="Verified Attack", exact=True)
    selected_tag.wait_for(state="visible", timeout=5000)

    print("Verified Attack selected")

    #search
    search_btn = page.locator('.search-btn[ng-click="searchEssential()"]:visible')
    print("Search btn found: ", search_btn.count())
    # Start listening before clicking so we do not miss a fast response.
    with page.expect_response(
        lambda response: urlsplit(response.url).path.endswith("/ddos_event_search_v2")
        and parse_qs(urlsplit(response.url).query).get("tag_key") == ["verified_attack"],
        timeout=30000,
    ) as search_response:
        search_btn.click()
    response = search_response.value
    response.body()
    if not response.ok:
        raise RuntimeError(f"Verified Attack search failed: HTTP {response.status}")

    #copy the attacks on every page
    all_attack_texts = []
    seen_ids = set()
    attacks = page.locator(".event-list-container:visible .nx-ddos-event-item")
    next_button = page.locator('li[ng-click="nextPage()"]:visible')
    pagination = next_button.locator("..")
    current_button = pagination.locator("li.active")
    number_buttons = pagination.locator('li[ng-repeat="item in pageList track by $index"]')

    while True:
        attacks.first.wait_for(state="visible", timeout=30000)
        next_button.wait_for(state="visible", timeout=10000)
        current_page = int(current_button.inner_text().strip())
        page_labels = number_buttons.all_inner_texts()
        last_page = max(
            int(label.strip())
            for label in page_labels
            if label.strip().isdigit()
        )

        # On other pages, disabled can mean the results are still loading.
        if current_page < last_page:
            expect(next_button).not_to_have_class(re.compile(r"\bdisabled\b"), timeout=30000)

        page_texts = attacks.all_inner_texts()
        for text in page_texts:
            alert_id = parse_attack(text)["alert_id"]
            if alert_id not in seen_ids:
                all_attack_texts.append(text)
                seen_ids.add(alert_id)
        print(f"Page {current_page}/{last_page}: {len(page_texts)} attacks "
              f"({len(all_attack_texts)} unique attacks collected)")

        # Collect the last page before exiting the loop.
        if current_page == last_page:
            break

        previous_id = parse_attack(page_texts[0])["alert_id"]
        with page.expect_response(
            lambda response: urlsplit(response.url).path.endswith("/ddos_event_search_v2"),
            timeout=30000,
        ) as next_response:
            next_button.click()
        response = next_response.value
        response.body()
        if not response.ok:
            raise RuntimeError(f"Page {current_page + 1} search failed: HTTP {response.status}")

        # The HTTP response can arrive before the new cards are rendered.
        expect(current_button).to_have_text(str(current_page + 1), timeout=30000)
        expect(attacks.first).not_to_contain_text(previous_id, timeout=30000)
        attacks.first.wait_for(state="visible", timeout=30000)

    raw_file = Path(__file__).resolve().with_name("attacks_raw.json")
    with open(raw_file, "w", encoding="utf-8") as file:
        json.dump(all_attack_texts, file, ensure_ascii=False, indent=2)

    parsed_attacks = []
    for text in all_attack_texts:
        attack = parse_attack(text)
        parsed_attacks.append(attack)
    parsed_file = Path(__file__).resolve().with_name("parsed_attack.json")
    with open(parsed_file, "w", encoding="utf-8") as file:
        json.dump(parsed_attacks, file, ensure_ascii=False, indent=2)
    print(f"Saved {len(parsed_attacks)} parsed attacks to {parsed_file}")

    return page



def create_excel_file(attack_file, output_path, year, template_path=None):
    """Write one row per JSON attack using the supplied ddos.xlsx layout."""
    with open(attack_file, encoding="utf-8") as file:
        alerts = json.load(file)
    if not isinstance(alerts, list) or any(not isinstance(alert, dict) for alert in alerts):
        raise ValueError("Expected a JSON list of parsed attack dictionaries")

    alerts.sort(key=lambda alert: datetime.strptime(
        f"{year} {alert['start_time']}", "%Y %b %d %H:%M:%S"
    ))

    template_path = Path(template_path) if template_path else Path(__file__).with_name("ddos.xlsx")
    output_path = Path(output_path)
    if output_path.resolve() in (template_path.resolve(), Path(attack_file).resolve()):
        raise ValueError("Output must be different from the template and JSON input")
    wb = openpyxl.load_workbook(template_path)
    sheet = wb.active
    # Keep the template header, widths and cell styles, but remove example attacks.
    row_styles = [copy(cell._style) for cell in sheet[2]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    months = {name: number for number, name in enumerate(
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), 1)}
    factors = {"": 1, "K": 1000, "M": 1000000, "G": 1000000000, "T": 1000000000000}

    for row, alert in enumerate(alerts, start=2):
        start = alert.get("start_time") or ""
        end = alert.get("end_time") or "(ongoing)"
        month = months.get(start.split()[0]) if start else None
        if month is None:
            raise ValueError(f"Row {row}: cannot read month from {start!r}")

        # Convert Mbps/Gbps and Kpps/Mpps before comparing the thresholds.
        bandwidth = alert.get("bandwidth")
        packet_rate = alert.get("packet_rate")
        gbps = mpps = None
        if bandwidth:
            match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([KMGT]?)bps", bandwidth)
            if not match:
                raise ValueError(f"Row {row}: invalid bandwidth {bandwidth!r}")
            gbps = float(match.group(1)) * factors[match.group(2)] / 1000000000
        if packet_rate:
            match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([KMGT]?)pps", packet_rate)
            if not match:
                raise ValueError(f"Row {row}: invalid packet rate {packet_rate!r}")
            mpps = float(match.group(1)) * factors[match.group(2)] / 1000000

        values = {
            "A": month,
            "B": year,  # Nexusguard's displayed timestamps do not include a year.
            "C": f"{start} - {end}",
            "D": alert.get("duration"),
            "G": "x" if mpps is not None and mpps > 1 else None,
            "J": f"{bandwidth or 'N/A'} / {packet_rate or 'N/A'}",
            "K": alert.get("ip"),
            "L": alert.get("profile"),
            "Q": gbps,
        }
        for column, value in values.items():
            cell = sheet[f"{column}{row}"]
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # Imported text must not become an Excel formula.

        # These independent flags recalculate if the numeric Gbps value changes.
        for column, condition in {
            "E": f"Q{row}>10", "F": f"Q{row}>40",
            "H": f"AND(Q{row}>1,Q{row}<10)", "I": f"Q{row}>100",
        }.items():
            sheet[f"{column}{row}"] = f'=IF(Q{row}="","",IF({condition},"x",""))'

        # M is the template spacer. N/O/P/R need data not present in this JSON.
        for column in range(1, 19):
            cell = sheet.cell(row, column)
            cell._style = copy(row_styles[column - 1])
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet[f"Q{row}"].number_format = '0.#########'
        sheet.row_dimensions[row].height = 42

    for cell in sheet[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 58
    for column in "EFGHI":
        sheet.column_dimensions[column].width = 13
    sheet.column_dimensions["C"].width = 42
    sheet.column_dimensions["K"].width = 23
    sheet.column_dimensions["L"].width = 36
    sheet.column_dimensions["O"].width = 14
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:R{max(1, len(alerts) + 1)}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    print(f"Saved {len(alerts)} attacks to {output_path}")
    return output_path



def main():
    with sync_playwright() as p:
        #open the browser, search for verified attack
        browser = p.chromium.launch()
        page = search_verified_attack(browser)

        #generate the excel file
        create_excel_file(
            Path(__file__).with_name("parsed_attack.json"),
            Path(__file__).parent / "outputs" / "nexusguard" / "verified_attacks.xlsx",
            year=2026,
        )

        input("Press Enter to close browser...")
        browser.close()
        os.unlink("attacks_raw.json")
        os.unlink("parsed_attacks.json")

if __name__ == "__main__":
    main()
